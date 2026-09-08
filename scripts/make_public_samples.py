"""Generate every sample asset the public topdf pages hand out.

Separate from `make_samples.py`, which builds the *client* tool's samples and
imitates that client's catalog design. Nothing is shared between them: different
palette, different products, different template layout, different columns. That
separation is the point — the public product must not carry the client's look.

    ./.venv/bin/python scripts/make_public_samples.py

Writes into web/static/samples/. Safe to re-run; it overwrites in place.

It verifies as it goes: every template it draws is run back through
`simple_catalog.find_product_boxes`, and the script fails rather than shipping a
template the tool cannot read.
"""

from __future__ import annotations

import csv
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from PIL import Image, ImageDraw
from reportlab.lib.colors import HexColor
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas as pdf_canvas

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO))

from catalog_exporter import render_pdf_page, try_font  # noqa: E402
from simple_catalog import export_catalog, find_product_boxes  # noqa: E402

SAMPLES = REPO / "web" / "static" / "samples"
PHOTOS = SAMPLES / "photos"

PAGE_W, PAGE_H = A4

# A palette of its own. Deliberately nothing like the client catalog's red.
INK = "#1F2933"
BRAND = "#0F5257"
ACCENT = "#2E8B93"
BOX_LINE = "#8A9199"
BOX_LINE_WIDTH = 1.6
MUTED = "#7B8794"

# The public sample images are fetched over the network by the exporter, the
# same way a real user's would be, so they need an address that resolves.
PHOTO_BASE = "https://topdf.stuffs.bid/topdf/static/samples/photos"

# Workshop and site supplies: a different trade from the client's grocery
# catalog, and a good fit for "anyone who sells physical goods".
SAMPLE_PRODUCTS = [
    ("Hex Bolt M8 x 40", "Zinc-plated steel, DIN 933", "0.42", "200", "bolt", "#5B6B7B"),
    ("Nylon Lock Nut M8", "Nyloc, zinc-plated, DIN 985", "0.18", "500", "nut", "#79808A"),
    ("Nitrile Gloves, Blue", "Powder-free, 4 mil, size L", "6.90", "100", "box", "#2E6FA7"),
    ("PTFE Thread Tape", "12mm x 12m, 0.075mm gauge", "0.85", "50", "roll", "#C9CDD2"),
    ("Cable Ties 200mm", "UV-stable black nylon", "3.40", "1000", "bag", "#3C4450"),
    ("Masking Tape 48mm", "Low-tack, 50m roll", "1.95", "36", "roll", "#D8B25A"),
    ("Safety Glasses, Clear", "Anti-fog, EN166 rated", "2.75", "24", "glasses", "#6BA3C4"),
    ("Abrasive Flap Disc 115mm", "80 grit zirconium", "1.60", "40", "disc", "#8C5A3C"),
    ("Copper Pipe Clip 15mm", "Single-fix, brass finish", "0.22", "250", "clip", "#B07A4B"),
    ("Silicone Sealant, Clear", "310ml cartridge, neutral cure", "3.10", "25", "tube", "#9AA5B1"),
    ("Wire Wheel Brush 75mm", "Crimped steel, 6mm shank", "2.40", "20", "disc", "#6E7681"),
    ("Utility Knife Blades", "0.5mm carbon steel, 100 pack", "4.20", "20", "box", "#4A5563"),
]


# --------------------------------------------------------------------------
# Product photos
# --------------------------------------------------------------------------

def photo_filename(index: int) -> str:
    return f"item-{index + 1:02d}.png"


def draw_photo(shape: str, colour: str, path: Path) -> None:
    """A flat illustration of a workshop item, on white.

    White background is deliberate: the exporter trims white before placing the
    image, which is what a real product cut-out gets too.
    """
    width, height = 640, 640
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    fill = HexColor(colour)
    body = (int(fill.red * 255), int(fill.green * 255), int(fill.blue * 255))
    shade = tuple(max(0, channel - 42) for channel in body)
    light = tuple(min(255, channel + 34) for channel in body)
    cx = cy = width // 2

    if shape == "bolt":
        draw.regular_polygon((cx, 190, 118), n_sides=6, rotation=30, fill=body)
        draw.regular_polygon((cx, 190, 78), n_sides=6, rotation=30, fill=shade)
        draw.rectangle((cx - 52, 250, cx + 52, 520), fill=body)
        for y in range(270, 520, 26):
            draw.line((cx - 52, y, cx + 52, y - 12), fill=shade, width=7)
    elif shape == "nut":
        draw.regular_polygon((cx, cy, 190), n_sides=6, rotation=30, fill=body)
        draw.ellipse((cx - 88, cy - 88, cx + 88, cy + 88), fill="white")
        draw.ellipse((cx - 88, cy - 88, cx + 88, cy + 88), outline=shade, width=10)
    elif shape == "box":
        draw.polygon([(cx - 200, 250), (cx + 60, 175), (cx + 210, 245), (cx - 50, 322)], fill=light)
        draw.rectangle((cx - 200, 250, cx - 50, 500), fill=body)
        draw.polygon([(cx - 50, 322), (cx + 210, 245), (cx + 210, 425), (cx - 50, 500)], fill=shade)
    elif shape == "roll":
        draw.ellipse((cx - 175, 170, cx + 175, 300), fill=light)
        draw.rectangle((cx - 175, 235, cx + 175, 440), fill=body)
        draw.ellipse((cx - 175, 375, cx + 175, 505), fill=shade)
        draw.ellipse((cx - 62, 205, cx + 62, 265), fill="white")
    elif shape == "bag":
        draw.polygon([(cx - 185, 250), (cx + 185, 250), (cx + 150, 520), (cx - 150, 520)], fill=body)
        draw.polygon([(cx - 185, 250), (cx - 130, 165), (cx + 130, 165), (cx + 185, 250)], fill=shade)
    elif shape == "glasses":
        draw.rounded_rectangle((cx - 210, 265, cx - 20, 385), radius=44, fill=light, outline=shade, width=8)
        draw.rounded_rectangle((cx + 20, 265, cx + 210, 385), radius=44, fill=light, outline=shade, width=8)
        draw.line((cx - 20, 300, cx + 20, 300), fill=shade, width=14)
    elif shape == "disc":
        draw.ellipse((cx - 195, cy - 195, cx + 195, cy + 195), fill=body)
        draw.ellipse((cx - 195, cy - 195, cx + 195, cy + 195), outline=shade, width=12)
        draw.ellipse((cx - 48, cy - 48, cx + 48, cy + 48), fill="white", outline=shade, width=8)
    elif shape == "clip":
        draw.arc((cx - 175, cy - 160, cx + 175, cy + 190), start=200, end=340, fill=body, width=52)
        draw.rectangle((cx - 165, cy + 120, cx - 65, cy + 190), fill=shade)
        draw.rectangle((cx + 65, cy + 120, cx + 165, cy + 190), fill=shade)
    else:  # tube
        draw.rounded_rectangle((cx - 88, 205, cx + 88, 520), radius=26, fill=body)
        draw.polygon([(cx - 46, 205), (cx + 46, 205), (cx + 26, 120), (cx - 26, 120)], fill=shade)
        draw.rectangle((cx - 30, 92, cx + 30, 130), fill=light)

    image.save(path, "PNG", optimize=True)


# --------------------------------------------------------------------------
# The product file
# --------------------------------------------------------------------------

HEADERS = ["Image URL", "Name", "Description", "Price", "Case Size"]


def sample_rows() -> list[list[str]]:
    return [
        [f"{PHOTO_BASE}/{photo_filename(index)}", name, description, price, case]
        for index, (name, description, price, case, _shape, _colour) in enumerate(SAMPLE_PRODUCTS)
    ]


def write_xlsx(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Products"
    sheet.append(HEADERS)
    for row in sample_rows():
        sheet.append(row)

    header_fill = PatternFill("solid", fgColor="0F5257")
    for column, _ in enumerate(HEADERS, start=1):
        cell = sheet.cell(row=1, column=column)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="left")
    for column, width in enumerate((58, 26, 34, 10, 12), start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = "A2"
    workbook.save(path)


def write_csv(path: Path) -> None:
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(HEADERS)
        writer.writerows(sample_rows())


# --------------------------------------------------------------------------
# The template PDFs
# --------------------------------------------------------------------------

def draw_header(page: pdf_canvas.Canvas, subtitle: str) -> None:
    page.setFillColor(HexColor(BRAND))
    page.rect(0, PAGE_H - 74, PAGE_W, 74, stroke=0, fill=1)
    page.setFillColor(HexColor("#FFFFFF"))
    page.setFont("Helvetica-Bold", 17)
    page.drawString(42, PAGE_H - 46, "YOUR LOGO HERE")
    page.setFont("Helvetica", 9.5)
    page.setFillColor(HexColor("#BEE0E2"))
    page.drawRightString(PAGE_W - 42, PAGE_H - 45, subtitle)


def draw_footer(page: pdf_canvas.Canvas, text: str) -> None:
    page.setFillColor(HexColor(MUTED))
    page.setFont("Helvetica", 7.5)
    page.drawCentredString(PAGE_W / 2, 24, text)


def draw_boxes(page: pdf_canvas.Canvas, rows: int, top: float, bottom: float) -> None:
    """A `rows` x 3 grid of empty rounded boxes between two y positions."""
    margin = 42.0
    gutter = 14.0
    columns = 3
    box_w = (PAGE_W - 2 * margin - (columns - 1) * gutter) / columns
    box_h = (top - bottom - (rows - 1) * gutter) / rows

    page.setStrokeColor(HexColor(BOX_LINE))
    page.setLineWidth(BOX_LINE_WIDTH)
    for row in range(rows):
        y = top - (row + 1) * box_h - row * gutter
        for column in range(columns):
            x = margin + column * (box_w + gutter)
            page.roundRect(x, y, box_w, box_h, 6, stroke=1, fill=0)


def write_one_page_template(path: Path) -> int:
    page = pdf_canvas.Canvas(str(path), pagesize=A4)
    draw_header(page, "Product catalogue")
    draw_boxes(page, rows=4, top=PAGE_H - 96, bottom=52)
    draw_footer(page, "Sample one-page template — this page repeats for the whole catalogue")
    page.showPage()
    page.save()
    return 12


def write_two_page_template(path: Path) -> tuple[int, int]:
    page = pdf_canvas.Canvas(str(path), pagesize=A4)

    # Page 1: a cover with room for a title, so fewer boxes.
    draw_header(page, "Spring price list")
    page.setFillColor(HexColor(INK))
    page.setFont("Helvetica-Bold", 24)
    page.drawString(42, PAGE_H - 126, "Catalogue title goes here")
    page.setFillColor(HexColor(MUTED))
    page.setFont("Helvetica", 10.5)
    page.drawString(42, PAGE_H - 146, "Replace this line with your own strapline, season or issue number.")
    draw_boxes(page, rows=3, top=PAGE_H - 172, bottom=52)
    draw_footer(page, "Sample template, page 1 — the cover")
    page.showPage()

    # Page 2: the denser inner page, which repeats.
    draw_header(page, "Spring price list")
    draw_boxes(page, rows=4, top=PAGE_H - 96, bottom=52)
    draw_footer(page, "Sample template, page 2 — repeats until the products run out")
    page.showPage()

    page.save()
    return 9, 12


# --------------------------------------------------------------------------
# Verification and previews
# --------------------------------------------------------------------------

def verify_template(path: Path, expected: list[int]) -> None:
    """Fail here rather than shipping a template the tool cannot read."""
    for index, wanted in enumerate(expected, start=1):
        rendered = render_pdf_page(path, index, dpi=150)
        try:
            found = len(find_product_boxes(rendered))
        finally:
            rendered.close()
        if found != wanted:
            raise SystemExit(
                f"{path.name} page {index}: detector found {found} product boxes, expected {wanted}. "
                f"The template and the detector have to agree before this ships."
            )
        print(f"  verified {path.name} page {index}: {found} boxes")


def write_previews(pdf_path: Path, stem: str, pages: int, dpi: int = 76) -> None:
    for number in range(1, pages + 1):
        rendered = render_pdf_page(pdf_path, number, dpi=dpi)
        try:
            rendered.save(SAMPLES / f"{stem}-page{number}.png", "PNG", optimize=True)
        finally:
            rendered.close()


def write_output_preview(template: Path, stem: str, pages: int) -> None:
    output = SAMPLES / f"_{stem}.pdf"
    export_catalog(
        excel_path=SAMPLES / "products-template.xlsx",
        template_pdf=template,
        output_path=output,
        max_products=400,
    )
    write_previews(output, stem, pages)
    output.unlink(missing_ok=True)


def write_hero_shot() -> None:
    """The picture at the top of the landing page.

    Separate from the before/after samples on purpose. Those show the *upload*
    flow, so their brand bar reads "YOUR LOGO HERE" -- correct there, because
    the blank template beside them is the thing you download and put your own
    logo on. In the hero it was wrong twice over: the headline sells the
    Shopify flow, and a placeholder logo is the single clearest sign that a
    page is a demo rather than a product.

    So this one is built by `shopify_catalog`'s own template generator, with a
    brand profile standing in for a store's. The hero then shows exactly what
    the paste-a-URL path produces -- including that the shop name is typeset
    when no logo image is available, which is a real outcome of that path.

    The shop is invented. Publishing a real store's catalog as our own
    marketing is not ours to do, however good it looks.
    """
    import shopify_catalog
    from simple_catalog import export_catalog as export_public

    brand = shopify_catalog.BrandProfile(
        name="Northgate Supply Co.",
        tagline="Workshop and site consumables, trade only",
        address="Sheffield, United Kingdom",
        primary="#123B47",
        secondary="#2E8B93",
    )

    template = SAMPLES / "_hero-template.pdf"
    output = SAMPLES / "_hero-export.pdf"
    boxes, sources = shopify_catalog.build_template(
        brand, "northgatesupply.example", template, len(SAMPLE_PRODUCTS)
    )
    export_public(
        excel_path=SAMPLES / "products-template.xlsx",
        template_pdf=template,
        output_path=output,
        max_products=400,
        known_boxes=boxes,
        page_sources=sources,
    )
    rendered = render_pdf_page(output, 1, dpi=112)
    try:
        rendered.save(SAMPLES / "hero-export.png", "PNG", optimize=True)
    finally:
        rendered.close()
    template.unlink(missing_ok=True)
    output.unlink(missing_ok=True)
    print("  wrote hero-export.png")


def main() -> None:
    SAMPLES.mkdir(parents=True, exist_ok=True)
    PHOTOS.mkdir(parents=True, exist_ok=True)

    print("Drawing product photos...")
    for index, (_name, _description, _price, _case, shape, colour) in enumerate(SAMPLE_PRODUCTS):
        draw_photo(shape, colour, PHOTOS / photo_filename(index))

    print("Writing the product file...")
    write_xlsx(SAMPLES / "products-template.xlsx")
    write_csv(SAMPLES / "products-template.csv")

    print("Drawing templates...")
    one_page = SAMPLES / "template-one-page.pdf"
    two_page = SAMPLES / "template-two-page.pdf"
    expected_one = write_one_page_template(one_page)
    expected_two = write_two_page_template(two_page)

    print("Verifying the detector agrees with them...")
    verify_template(one_page, [expected_one])
    verify_template(two_page, list(expected_two))

    print("Rendering previews...")
    write_previews(one_page, "template-one-page", 1)
    write_previews(two_page, "template-two-page", 2)

    print("Exporting the samples for real...")
    write_output_preview(one_page, "example-one-page", 1)
    write_output_preview(two_page, "example-two-page", 2)

    print("Building the hero shot from the Shopify template generator...")
    write_hero_shot()

    print("Done.")


if __name__ == "__main__":
    main()
